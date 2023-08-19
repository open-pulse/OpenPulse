from PyQt5.QtWidgets import *
from PyQt5.QtGui import *
from PyQt5.QtCore import Qt
from PyQt5 import uic
from pathlib import Path

import os
import numpy as np
import pandas as pd
import openpyxl
import matplotlib.pyplot as plt

from pulse.tools.advanced_cursor import AdvancedCursor
from pulse.postprocessing.plot_structural_data import get_structural_frf
from data.user_input.project.printMessageInput import PrintMessageInput

def get_icons_path(filename):
    path = f"data/icons/{filename}"
    if os.path.exists(path):
        return str(Path(path))

window_title1 = "ERROR MESSAGE"
window_title2 = "WARNING MESSAGE"


class PlotStructuralFrequencyResponseInput(QDialog):
    def __init__(self, project, opv, analysisMethod, solution, *args, **kwargs):
        super().__init__(*args, **kwargs)

        uic.loadUi(Path('data/user_input/ui/plots_/results_/structural_/plotStructuralFrequencyResponseInput.ui'), self)

        self.icon = QIcon(get_icons_path('pulse.png'))
        self.search_icon = QIcon(get_icons_path('searchFile.png'))
        self.setWindowIcon(self.icon)

        self.setWindowFlags(Qt.WindowStaysOnTopHint)
        self.setWindowModality(Qt.WindowModal)

        self.opv = opv
        self.opv.setInputObject(self)
        self.list_node_IDs = self.opv.getListPickedPoints()

        self.projec = project
        self.preprocessor = project.preprocessor
        self.before_run = project.get_pre_solution_model_checks()
        self.nodes = self.preprocessor.nodes
        
        self.analysisMethod = analysisMethod
        self.frequencies = project.frequencies
        self.solution = solution

        self.initialize_variables()
        self.define_and_configure_Qt_variables()
        self.create_connections()
        self.update_skiprows_visibility()
        self.writeNodes(self.list_node_IDs)
        self.exec()

    def initialize_variables(self):
        """
        """
        self.userPath = os.path.expanduser('~')
        self.imported_path = ""
        self.save_path = ""
        self.node_ID = 0
        self.imported_data = None
        self.localDof = None
        self.imported_results = {}
        self.ids_to_checkBox = {}
        self.checkButtons_state = {}

    def define_and_configure_Qt_variables(self):
        """
        """
        # CheckBox
        self.checkBox_cursor = self.findChild(QCheckBox, 'checkBox_cursor')
        self.checkBox_skiprows = self.findChild(QCheckBox, "checkBox_skiprows")
        
        # LineEdit
        self.lineEdit_nodeID = self.findChild(QLineEdit, 'lineEdit_nodeID')
        self.lineEdit_FileName = self.findChild(QLineEdit, 'lineEdit_FileName')
        self.lineEdit_ImportResultsPath = self.findChild(QLineEdit, 'lineEdit_ImportResultsPath')
        self.lineEdit_SaveResultsPath = self.findChild(QLineEdit, 'lineEdit_SaveResultsPath')
        self.lineEdit_ImportResultsPath.setDisabled(True)

        # PushButton
        self.pushButton_plot_frequency_response = self.findChild(QPushButton, 'pushButton_plot_frequency_response')
        self.pushButton_search_file_to_import = self.findChild(QPushButton, 'pushButton_search_file_to_import')
        self.pushButton_search_file_to_import.setIcon(self.search_icon)

        # RadioButton
        self.radioButton_ux = self.findChild(QRadioButton, 'radioButton_ux')
        self.radioButton_uy = self.findChild(QRadioButton, 'radioButton_uy')
        self.radioButton_uz = self.findChild(QRadioButton, 'radioButton_uz')
        self.radioButton_rx = self.findChild(QRadioButton, 'radioButton_rx')
        self.radioButton_ry = self.findChild(QRadioButton, 'radioButton_ry')
        self.radioButton_rz = self.findChild(QRadioButton, 'radioButton_rz')
        self.radioButton_plotAbs = self.findChild(QRadioButton, 'radioButton_plotAbs')
        self.radioButton_plotReal = self.findChild(QRadioButton, 'radioButton_plotReal')
        self.radioButton_plotImag = self.findChild(QRadioButton, 'radioButton_plotImag')
        self.radioButton_Absolute = self.findChild(QRadioButton, 'radioButton_Absolute')
        self.radioButton_Real_Imaginary = self.findChild(QRadioButton, 'radioButton_Real_Imaginary')
        self.radioButton_NoneDiff = self.findChild(QRadioButton, 'radioButton_NoneDiff')
        self.radioButton_SingleDiff = self.findChild(QRadioButton, 'radioButton_SingleDiff')
        self.radioButton_DoubleDiff = self.findChild(QRadioButton, 'radioButton_DoubleDiff')

        # SpinBox
        self.spinBox_skiprows = self.findChild(QSpinBox, 'spinBox')

        # TabWidget
        self.tabWidget_plot_results = self.findChild(QTabWidget, "tabWidget_plot_results")
        self.tab_plot = self.tabWidget_plot_results.findChild(QWidget, "tab_plot")

        # ToolButton
        self.toolButton_ChooseFolderExport = self.findChild(QToolButton, 'toolButton_ChooseFolderExport')
        self.toolButton_ExportResults = self.findChild(QToolButton, 'toolButton_ExportResults')
        self.toolButton_ResetPlot = self.findChild(QToolButton, 'toolButton_ResetPlot')
        
        # TreeWidget
        self.treeWidget_import_text_files = self.findChild(QTreeWidget, "treeWidget_import_text_files")
        self.treeWidget_import_sheet_files = self.findChild(QTreeWidget, "treeWidget_import_sheet_files")

        widths_1 = [320, 60]
        for i, width in enumerate(widths_1):
            self.treeWidget_import_text_files.setColumnWidth(i, width)

        widths_2 = [180, 150, 60]
        for i, width in enumerate(widths_2):
            self.treeWidget_import_sheet_files.setColumnWidth(i, width)

    def create_connections(self):
        """
        """
        self.pushButton_search_file_to_import.clicked.connect(self.choose_path_import_results)
        self.toolButton_ChooseFolderExport.clicked.connect(self.choose_path_export_results)
        self.toolButton_ExportResults.clicked.connect(self.ExportResults)
        self.toolButton_ResetPlot.clicked.connect(self.reset_imported_data)
        self.use_cursor = self.checkBox_cursor.isChecked()
        self.checkBox_cursor.clicked.connect(self.update_cursor)
        self.checkBox_skiprows.clicked.connect(self.update_skiprows_visibility)

        self.Ux = self.radioButton_ux.isChecked()
        self.Uy = self.radioButton_uy.isChecked()
        self.Uz = self.radioButton_uz.isChecked()
        self.Rx = self.radioButton_rx.isChecked()
        self.Ry = self.radioButton_ry.isChecked()
        self.Rz = self.radioButton_rz.isChecked()

        self.radioButton_Absolute.clicked.connect(self.radioButtonEvent_save_data)
        self.radioButton_Real_Imaginary.clicked.connect(self.radioButtonEvent_save_data)
        self.save_Absolute = self.radioButton_Absolute.isChecked()
        self.save_Real_Imaginary = self.radioButton_Real_Imaginary.isChecked()

        # self.pushButton_AddImportedPlot.clicked.connect(self.ImportResults)
        self.pushButton_plot_frequency_response.clicked.connect(self.check_inputs_and_plot)

    def update_cursor(self):
        self.use_cursor = self.checkBox_cursor.isChecked()

    def update_skiprows_visibility(self):
        self.spinBox_skiprows.setDisabled(not self.checkBox_skiprows.isChecked())

    def reset_imported_data(self):
        self.initialize_variables()
        self.update_treeWidget_info()
        self.lineEdit_ImportResultsPath.setText("")
        title = "Information"
        message = "The plot data has been reseted."
        PrintMessageInput([title, message, window_title2])
    
    def writeNodes(self, list_node_ids):
        text = ""
        for node in list_node_ids:
            text += "{}, ".format(node)
        self.lineEdit_nodeID.setText(text)

    def update(self):
        self.list_node_IDs = self.opv.getListPickedPoints()
        if self.list_node_IDs != []:
            self.writeNodes(self.list_node_IDs)

    def keyPressEvent(self, event):
        if event.key() == Qt.Key_Enter or event.key() == Qt.Key_Return:
            self.check_inputs_and_plot()
        elif event.key() == Qt.Key_Escape:
            self.close()

    def radioButtonEvent_save_data(self):
        """
        """
        self.save_Absolute = self.radioButton_Absolute.isChecked()
        self.save_Real_Imaginary = self.radioButton_Real_Imaginary.isChecked()

    def choose_path_import_results(self):
        """
        """
        if self.imported_path == "":
            _path = self.userPath
        else:
            _path = os.path.dirname(self.imported_path)

        self.imported_path, _ = QFileDialog.getOpenFileName(None, 'Open file', _path, 'Files (*.csv; *.dat; *.txt; *.xlsx; *.xls)')
        self.import_name = os.path.basename(self.imported_path)
        self.lineEdit_ImportResultsPath.setText(self.imported_path)
        
        if self.imported_path != "":
            if os.path.exists(self.imported_path):
                self.ImportResults()
                self.update_treeWidget_info()
    
    def get_data_index(self):
        """
        """
        index = 1
        run = True
        while run:
            if index not in self.imported_results.keys():
                key = index
                run = False
            else:
                index += 1
        return key

    def ImportResults(self):
        """
        """
        try:
            message = ""
            run = True
            if self.checkBox_skiprows.isChecked():
                skiprows = int(self.spinBox_skiprows.text())
            else:
                skiprows = 0
            maximum_lines_to_skip = 100
            
            while run:
                try:
                    sufix = Path(self.imported_path).suffix
                    filename = os.path.basename(self.imported_path)
                    if sufix in [".txt", ".dat", ".csv"]:
                        loaded_data = np.loadtxt(self.imported_path, delimiter=",", skiprows=skiprows)
                        key = self.get_data_index()
                        self.imported_results[key] = {  "data" : loaded_data,
                                                        "filename" : filename,
                                                        "extension" : sufix  }

                    elif sufix in [".xls", ".xlsx"]:
                        wb = openpyxl.load_workbook(self.imported_path)
                        sheetnames = wb.sheetnames
                        for sheetname in sheetnames:
                            sheet_data = pd.read_excel(self.imported_path, sheet_name=sheetname, header=skiprows, usecols=[0,1,2]).to_numpy()
                            key = self.get_data_index()
                            self.imported_results[key] = {  "data" : sheet_data,
                                                            "filename" : filename,
                                                            "sheetname" : sheetname,
                                                            "extension" : sufix  }

                    self.spinBox_skiprows.setValue(int(skiprows))
                    run = False

                except:
                    skiprows += 1
                    if skiprows>=maximum_lines_to_skip:
                        run = False
                        title = "Error while loading data from file"
                        message = "The maximum number of rows to skip has been reached and no valid data has "
                        message += "been found. Please, verify the data in the imported file to proceed."
                        message += "Maximum number of header rows: 100"

            # if skiprows<maximum_lines_to_skip:
            #     self.legend_imported = "imported data: "+ os.path.basename(self.imported_path).split(".")[0]
            #     self.tabWidget_plot_results.setCurrentWidget(self.tab_plot)
            #     title = "Information"
            #     message = "The results have been imported."
            #     PrintMessageInput([title, message, window_title2])
            #     return

        except Exception as log_error:
            title = "Error while loading data from file"
            message = str(log_error)
            return
        
        if message != "":
            PrintMessageInput([title, message, window_title1])


    def update_treeWidget_info(self):
        """
        """
        self.cache_checkButtons_state()
        self.treeWidget_import_text_files.clear()
        self.treeWidget_import_sheet_files.clear()
        #
        if len(self.imported_results) > 0:
            for i, (id, data) in enumerate(self.imported_results.items()):
                # Creates the QCheckButtons to control data to be plotted
                self.ids_to_checkBox[id] = QCheckBox()
                self.ids_to_checkBox[id].setStyleSheet("margin-left:30%; margin-right:50%;")

                if id in self.checkButtons_state.keys():
                    self.ids_to_checkBox[id].setChecked(self.checkButtons_state[id])

                if "sheetname" in data.keys():
                    _item = QTreeWidgetItem([str(data["filename"]), str(data["sheetname"])])
                    self.treeWidget_import_sheet_files.addTopLevelItem(_item)
                    self.treeWidget_import_sheet_files.setItemWidget(_item, 2, self.ids_to_checkBox[id])
                else:
                    _item = QTreeWidgetItem([str(data["filename"])])
                    self.treeWidget_import_text_files.addTopLevelItem(_item)
                    self.treeWidget_import_text_files.setItemWidget(_item, 1, self.ids_to_checkBox[id])                  

                for i in range(5):
                    _item.setTextAlignment(i, Qt.AlignCenter)
            

    def cache_checkButtons_state(self):
        """
        """
        self.checkButtons_state = {}
        for key, check in self.ids_to_checkBox.items():
            self.checkButtons_state[key] = check.isChecked()


    def choose_path_export_results(self):
        """
        """
        self.save_path = QFileDialog.getExistingDirectory(None, 'Choose a folder to export the results', self.userPath)
        self.save_name = os.path.basename(self.save_path)
        self.lineEdit_SaveResultsPath.setText(str(self.save_path))


    def check_inputs_and_plot(self, export=False):
        """
        """
        lineEdit_nodeID = self.lineEdit_nodeID.text()
        stop, self.node_ID = self.before_run.check_input_NodeID(lineEdit_nodeID, single_ID=True)
        if stop:
            return True

        self.localDof = None
        if self.radioButton_SingleDiff.isChecked():
            _unit_label = "m/s"
        elif self.radioButton_DoubleDiff.isChecked():
            _unit_label = "m/s²"
        else:
            _unit_label = "m"    

        if self.radioButton_ux.isChecked():
            self.localDof = 0
            self.localdof_label = "Ux"
            self.unit_label = _unit_label

        if self.radioButton_uy.isChecked():
            self.localDof = 1
            self.localdof_label = "Uy"
            self.unit_label = _unit_label

        if self.radioButton_uz.isChecked():
            self.localDof = 2
            self.localdof_label = "Uz"
            self.unit_label = _unit_label
 
        if self.radioButton_rx.isChecked():
            self.localDof = 3
            self.localdof_label = "Rx"
            self.unit_label = _unit_label

        if self.radioButton_ry.isChecked():
            self.localDof = 4
            self.localdof_label = "Ry"
            self.unit_label = _unit_label

        if self.radioButton_rz.isChecked():
            self.localDof = 5
            self.localdof_label = "Rz"
            self.unit_label = _unit_label

        if self.radioButton_SingleDiff.isChecked():
            _unit_label = "rad/s"
        elif self.radioButton_DoubleDiff.isChecked():
            _unit_label = "rad/s²"
        else:
            _unit_label = "rad"

        if not export:
            self.plot()

        return False


    def ExportResults(self):
        
        if self.lineEdit_FileName.text() != "":
            if self.save_path != "":
                self.export_path_folder = self.save_path + "/"
            else:
                title = "None folder selected"
                message = "Plese, choose a folder before trying export the results."
                PrintMessageInput([title, message, window_title1])
                return
        else:
            title = "Empty file name"
            message = "Inform a file name before trying export the results."
            PrintMessageInput([title, message, window_title1])
            return
        
        if self.check_inputs_and_plot(export=True):
            return

        freq = self.frequencies
        self.export_path = self.export_path_folder + self.lineEdit_FileName.text() + ".dat"
        response = self.get_response()

        if self.save_Absolute:
            header = ("Frequency[Hz], Real part [{}], Imaginary part [{}], Absolute [{}]").format(self.unit_label, self.unit_label, self.unit_label)
            data_to_export = np.array([freq, np.real(response), np.imag(response), np.abs(response)]).T
        elif self.save_Real_Imaginary:
            header = ("Frequency[Hz], Real part [{}], Imaginary part [{}]").format(self.unit_label, self.unit_label)
            data_to_export = np.array([freq, np.real(response), np.imag(response)]).T        
            
        np.savetxt(self.export_path, data_to_export, delimiter=",", header=header)
        title = "Information"
        message = "The results have been exported."
        PrintMessageInput([title, message, window_title2])


    def get_response(self):
        response = get_structural_frf(self.preprocessor, self.solution, self.node_ID, self.localDof)
        if self.radioButton_SingleDiff.isChecked():
            output_data = response*(1j*2*np.pi)*self.frequencies
        elif self.radioButton_DoubleDiff.isChecked():
            output_data = response*((1j*2*np.pi*self.frequencies)**2)
        else:
            output_data = response
        return output_data


    def plot(self):
        """
        """
        plt.ion()
        self.fig = plt.figure(figsize=[10,8])
        self.ax = self.fig.add_subplot(1,1,1)

        frequencies = self.frequencies
        response = self.get_response()

        list_plots = []
        list_legends = []
        colors = [  [0,0,0], [0,0,1], [0,1,0], [1,1,0], [0,1,1],
                    [1,0,1], [0.75,0.75,0.75], [0.5, 0.5, 0.5], [0.25, 0.25, 0.25]  ]
        
        legend_label = "Response {} at node {}".format(self.localdof_label, self.node_ID)
        plot_model = self.call_plotter(frequencies, response)
        list_plots.append(plot_model)
        list_legends.append(legend_label)
        j = 0
        for id, checkBox in self.ids_to_checkBox.items():
            if checkBox.isChecked():

                if id < len(colors):
                    color = colors[j]
                    j += 1
                else:
                    color = np.random.randint(0,255,3)/255

                # legend_label = f"Signal #{id}"
                data = self.imported_results[id]["data"]
                x_values = data[:, 0]
                y_values = data[:, 1] + 1j*data[:, 2]

                if "sheetname" in self.imported_results[id].keys():
                    sheetname = self.imported_results[id]["sheetname"]
                    legend_label = f"{sheetname}"
                else:
                    legend_label = self.imported_results[id]["filename"]

                plot_imported = self.call_plotter(x_values, y_values, color=color, linestyle="--")
                list_plots.append(plot_imported)
                list_legends.append(legend_label)

        _legends = plt.legend(handles=list_plots, labels=list_legends)
        plt.gca().add_artist(_legends)

        if self.analysisMethod is None:
            title = "NODAL RESPONSE"
        else:
            title = f"STRUCTURAL FREQUENCY RESPONSE - {self.analysisMethod.upper()}"

        self.ax.set_title(title, fontsize = 12, fontweight = 'bold')
        self.ax.set_xlabel('Frequency [Hz]', fontsize = 12, fontweight = 'bold')

        self.cursor = AdvancedCursor(self.ax, frequencies, response, self.use_cursor)
        self.mouse_connection = self.fig.canvas.mpl_connect(s='motion_notify_event', func=self.cursor.mouse_move)

        self.fig.show()


    def call_plotter(self, x_values, y_values, color=[1, 0, 0], linestyle="-"):
        """
        """
        if self.radioButton_plotAbs.isChecked():
            y_values = np.abs(y_values)
            self.ax.set_ylabel(("Structural Response - Absolute [{}]").format(self.unit_label), fontsize = 14, fontweight = 'bold')

            if not float(0) in y_values:
                self.ax.set_yscale('log')
                        
        elif self.radioButton_plotReal.isChecked():
            y_values = np.real(y_values)
            self.ax.set_ylabel(("Structural Response - Real [{}]").format(self.unit_label), fontsize = 14, fontweight = 'bold')
        
        elif self.radioButton_plotImag.isChecked():
            y_values = np.imag(y_values)
            self.ax.set_ylabel(("Structural Response - Imaginary [{}]").format(self.unit_label), fontsize = 14, fontweight = 'bold')

        aux_bool = self.radioButton_plotReal.isChecked() + self.radioButton_plotImag.isChecked()
        if float(0) in y_values or aux_bool:
            if float(0) in y_values[1:] or aux_bool:
                _plot, = plt.plot(x_values, y_values, color=color, linewidth=2, linestyle=linestyle)
            else:
                _plot, = plt.semilogy(x_values[1:], y_values[1:], color=color, linewidth=2, linestyle=linestyle)
        else: 
            _plot, = plt.semilogy(x_values, y_values, color=color, linewidth=2, linestyle=linestyle)
        return _plot